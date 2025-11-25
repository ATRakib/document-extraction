import os
from PyPDF2 import PdfReader
from docx import Document
from PIL import Image
import pytesseract
import openpyxl

class DocumentParser:
    @staticmethod
    def extract_text(file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.pdf':
            return DocumentParser.extract_from_pdf(file_path)
        elif ext in ['.docx', '.doc']:
            return DocumentParser.extract_from_docx(file_path)
        elif ext in ['.xlsx', '.xls']:
            return DocumentParser.extract_from_excel(file_path)
        elif ext in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff']:
            return DocumentParser.extract_from_image(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    @staticmethod
    def extract_from_pdf(file_path: str) -> str:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()

    @staticmethod
    def extract_from_docx(file_path: str) -> str:
        doc = Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text.strip()

    @staticmethod
    def extract_from_excel(file_path: str) -> str:
        workbook = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in workbook:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
        return text.strip()

    @staticmethod
    def extract_from_image(file_path: str) -> str:
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image)
        return text.strip()